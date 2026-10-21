"""Build a new DOCX from the fixed proposal JSON or a legacy Markdown draft.

Usage: bundled-python build_docx.py --input proposal.md --output proposal.docx
Fixed template: --input proposal.json --output proposal.docx
                [--template template.json] [--visual-dir rendered/visuals]
The JSON route shares its ordered content plan with native Google Docs and HTML.
Supports headings (levels 1–4), paragraphs, bold/italic, links, flat lists,
pipe tables, > conclusion callouts, standalone local PNG/JPEG images and
<!-- pagebreak -->. The first # heading is the cover title; the next two
paragraphs are its subtitle and agency/date line. Subsequent ##/###/#### map to
the reference's main/secondary/detail headings. Use --style woodcort|antropik.
Uses the bundled marked parser and python-docx; never edits the reference.
"""
from __future__ import annotations

import argparse
from io import BytesIO
import json
import os
from pathlib import Path, PurePosixPath
import subprocess
import sys
from urllib.parse import unquote, urlsplit

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm, Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph
from PIL import Image


FONT = "Arial"  # Explicit portable substitute for the reference's Helvetica Neue.
PAGEBREAK = "<!-- pagebreak -->"
STYLES = {
    "woodcort": {"background": "F3F3F3", "body": "000000", "heading": "000000", "accent": "000000",
                 "sizes": (28, 20, 15), "spaces": ((18, 18), (18, 10), (18, 8)), "title": 37,
                 "subtitle": 21, "subtitle_color": "000000", "meta": "6A6A6A",
                 "header": "D9EAF7", "header_text": "000000", "border": "C8C8C8", "border_size": "5",
                 "first_column": None, "callout": "FFF2CC", "callout_border": "D6B656", "callout_left_only": False},
    "antropik": {"background": "EFEFEF", "body": "333B47", "heading": "1F2945", "accent": "A1633B",
                 "sizes": (30, 22, 16), "spaces": ((20, 30), (50, 20), (35, 16)), "title": 50,
                 "subtitle": 20, "subtitle_color": "5C6678", "meta": "A1633B",
                 "header": "1F2A44", "header_text": "FFFFFF", "border": "D9DDE3", "border_size": "6",
                 "first_column": "F7F1EA", "callout": "E3EFDC", "callout_border": "6AA84F", "callout_left_only": True},
}


def tokens(markdown: str) -> list[dict]:
    dependencies = Path(sys.executable).resolve().parent.parent
    node = dependencies / "node" / "bin" / ("node.exe" if os.name == "nt" else "node")
    marked = dependencies / "node" / "node_modules" / "marked" / "lib" / "marked.esm.js"
    if not node.is_file() or not marked.is_file():
        raise ValueError("Run with the bundled workspace Python; bundled Node.js/marked are required")
    code = ("import fs from 'node:fs'; import {pathToFileURL} from 'node:url';"
            "const {marked}=await import(pathToFileURL(process.argv[1]).href);"
            "process.stdout.write(JSON.stringify(marked.lexer(fs.readFileSync(0,'utf8'),{gfm:true})));" )
    result = subprocess.run([str(node), "--input-type=module", "-e", code, str(marked)],
                            input=markdown, text=True, encoding="utf-8", capture_output=True,
                            check=True, timeout=30)
    return json.loads(result.stdout)


def inline(paragraph, items: list[dict], *, bold=False, italic=False) -> None:
    for item in items:
        kind = item["type"]
        if kind in {"strong", "em"}:
            inline(paragraph, item["tokens"], bold=bold or kind == "strong", italic=italic or kind == "em")
        elif kind in {"text", "escape"}:
            if kind == "text" and ("**" in item["text"] or "![" in item["text"]):
                raise ValueError("Malformed bold or image markup; escape literal Markdown characters")
            run = paragraph.add_run(item["text"].replace("\n", " "))
            run.bold, run.italic = (True if bold else None), (True if italic else None)
        elif kind == "link":
            href = item["href"]
            if urlsplit(href).scheme not in {"http", "https", "mailto"}:
                raise ValueError(f"Only http/https/mailto source links are supported: {href}")
            hyperlink = OxmlElement("w:hyperlink")
            hyperlink.set(qn("r:id"), paragraph.part.relate_to(href, RT.HYPERLINK, is_external=True))
            label = Paragraph(OxmlElement("w:p"), paragraph._parent)
            inline(label, item["tokens"], bold=bold, italic=italic)
            for run in label.runs:
                run.font.color.rgb = RGBColor.from_string("1155CC")
                run.underline = True
                hyperlink.append(run._r)
            paragraph._p.append(hyperlink)
        elif kind == "br":
            paragraph.add_run().add_break()
        else:
            raise ValueError(f"Unsupported inline Markdown: {kind}")


def table_cells(line: str) -> list[str]:
    """Count cells before marked can silently discard extra pipe-table values."""
    parts, cell, escaped = [], [], False
    for char in line.strip():
        if char == "|" and not escaped:
            parts.append("".join(cell)); cell = []
        else:
            cell.append(char)
        escaped = char == "\\" and not escaped
    parts.append("".join(cell))
    if line.strip().startswith("|"):
        parts.pop(0)
    if line.rstrip().endswith("|") and parts[-1] == "":
        parts.pop()
    return parts


def add_table(document, item: dict, theme: dict) -> None:
    count = len(item["header"])
    if not 2 <= count <= 7:
        raise ValueError("Tables require 2–7 columns; move wide registers into calculations")
    if any(len(table_cells(line)) != count for line in item["raw"].splitlines() if line.strip()):
        raise ValueError("Malformed table: every row must have the same number of cells")
    rows = [item["header"], *item["rows"]]
    table = document.add_table(rows=0, cols=count)
    table.autofit = False
    usable = document.sections[0].page_width - document.sections[0].left_margin - document.sections[0].right_margin
    weights = ([1] * count if theme["first_column"] else
               [max(12, min(48, max(len(row[col]["text"]) for row in rows))) for col in range(count)])
    for column, weight in zip(table.columns, weights):
        column.width = int(usable * weight / sum(weights))
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement("w:" + edge)
        for key, value in (("val", "single"), ("sz", theme["border_size"]), ("color", theme["border"])):
            border.set(qn("w:" + key), value)
        borders.append(border)
    table._tbl.tblPr.append(borders)
    for index, source in enumerate(rows):
        row = table.add_row()
        no_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(no_split)
        if index == 0:
            row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for col, (cell, value) in enumerate(zip(row.cells, source)):
            cell.width = table.columns[col].width
            props = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for edge in ("top", "left", "bottom", "right"):
                margin = OxmlElement("w:" + edge)
                margin.set(qn("w:w"), "100"); margin.set(qn("w:type"), "dxa")
                margins.append(margin)
            props.append(margins)
            fill = theme["header"] if index == 0 else (theme["first_column"] if col == 0 else "FFFFFF")
            if fill:
                shade = OxmlElement("w:shd"); shade.set(qn("w:fill"), fill); props.append(shade)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.keep_with_next = index == 0
            paragraph.paragraph_format.space_after = Pt(3)
            paragraph.paragraph_format.line_spacing = 1.1
            if item["align"][col] == "center":
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif item["align"][col] == "right":
                paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            inline(paragraph, value["tokens"], bold=index == 0)
            for run in paragraph.runs:
                run.font.size = Pt(10)
                run.font.color.rgb = RGBColor.from_string(theme["header_text"] if index == 0 else theme["body"])
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_image(document, item: dict, root: Path) -> None:
    href = item["href"]
    parsed = urlsplit(href)
    decoded = unquote(parsed.path)
    local = PurePosixPath(decoded)
    if (parsed.scheme or parsed.netloc or parsed.query or parsed.fragment or not decoded
            or "\\" in decoded or ":" in decoded or local.is_absolute() or ".." in local.parts):
        raise ValueError(f"Image must be a relative file inside the Markdown directory: {href}")
    path = (root / Path(*local.parts)).resolve(strict=True)
    if not path.is_relative_to(root) or path.suffix.lower() not in {".png", ".jpg", ".jpeg"}:
        raise ValueError(f"Image must be a local PNG/JPEG inside the Markdown directory: {href}")
    with Image.open(path) as image:
        if image.format not in {"PNG", "JPEG"}:
            raise ValueError(f"Image content must be PNG/JPEG: {href}")
        width, height = image.size
        image.verify()
    section = document.sections[0]
    max_width = section.page_width - section.left_margin - section.right_margin
    size = min(max_width, Inches(7.2) * width / height)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = bool(item["text"])
    paragraph.add_run().add_picture(str(path), width=int(size))
    if item["text"]:
        caption = document.add_paragraph(item["text"])
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in caption.runs:
            run.font.size = Pt(9); run.font.color.rgb = RGBColor.from_string("666666")


def build(input_path: Path, output_path: Path, style: str = "woodcort") -> Path:
    if output_path.exists() or output_path.is_symlink():
        raise ValueError(f"Output must be a new file: {output_path}")
    if input_path.suffix.lower() != ".md" or output_path.suffix.lower() != ".docx":
        raise ValueError("Input must be .md and output must be .docx")
    source = input_path.resolve(strict=True)
    content = source.read_text(encoding="utf-8-sig")
    if not content.strip():
        raise ValueError("Markdown document is empty")
    if style not in STYLES:
        raise ValueError(f"Unknown reference style: {style}")
    theme = STYLES[style]
    document = Document()
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    background = OxmlElement("w:background"); background.set(qn("w:color"), theme["background"])
    document._element.insert(0, background)
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(0.5)
    for name, size in (("Normal", 12), ("Title", theme["title"]), *[(f"Heading {i}", value) for i, value in enumerate(theme["sizes"], 1)]):
        paragraph_style = document.styles[name]
        if name == "Normal" and style == "woodcort":
            size = 11.5
        paragraph_style.font.name, paragraph_style.font.size = FONT, Pt(size)
        paragraph_style.font.color.rgb = RGBColor.from_string(theme["body"] if name == "Normal" else theme["heading"])
        paragraph_style.paragraph_format.space_before = Pt(10)
        paragraph_style.paragraph_format.space_after = Pt(10)
        paragraph_style.paragraph_format.line_spacing = 1
        if name == "Normal" and style == "woodcort":
            paragraph_style.paragraph_format.space_before = Pt(0)
            paragraph_style.paragraph_format.space_after = Pt(8)
        if name != "Normal":
            paragraph_style.font.bold = True
            paragraph_style.paragraph_format.keep_with_next = True
        if name.startswith("Heading"):
            level = int(name[-1]) - 1
            before, after = theme["spaces"][level]
            paragraph_style.paragraph_format.space_before = Pt(before)
            paragraph_style.paragraph_format.space_after = Pt(after)
            if level == 2:
                paragraph_style.font.color.rgb = RGBColor.from_string(theme["accent"])
    document.styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for border in document.styles["Title"].element.xpath("./w:pPr/w:pBdr"):
        border.getparent().remove(border)
    document.styles["Title"].paragraph_format.space_before = Pt(70)
    document.styles["Title"].paragraph_format.space_after = Pt(30)
    cover_paragraphs = 0
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE"); footer._p.append(field)
    for item in tokens(content):
        kind = item["type"]
        if kind == "space":
            continue
        if kind == "html" and item["raw"].strip() == PAGEBREAK:
            document.add_page_break()
            cover_paragraphs = 0
        elif kind == "heading" and item["depth"] <= 4:
            heading_style = ("Title" if not document.paragraphs and item["depth"] == 1 else
                             f'Heading {max(1, item["depth"] - 1)}')
            inline(document.add_paragraph(style=heading_style), item["tokens"])
            cover_paragraphs = 2 if heading_style == "Title" else 0
        elif kind == "paragraph":
            if len(item["tokens"]) == 1 and item["tokens"][0]["type"] == "image":
                add_image(document, item["tokens"][0], source.parent)
            else:
                if any(line.lstrip().startswith("|") for line in item["raw"].splitlines()):
                    raise ValueError("Malformed pipe table; include a valid header separator")
                paragraph = document.add_paragraph()
                inline(paragraph, item["tokens"])
                if cover_paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in paragraph.runs:
                        run.font.size = Pt(theme["subtitle"] if cover_paragraphs == 2 else 11.5)
                        run.font.color.rgb = RGBColor.from_string(theme["subtitle_color"] if cover_paragraphs == 2 else theme["meta"])
                    cover_paragraphs -= 1
        elif kind == "table":
            add_table(document, item, theme)
        elif kind == "blockquote":
            for quote in item["tokens"]:
                if quote["type"] == "space":
                    continue
                if quote["type"] != "paragraph":
                    raise ValueError("Conclusion callouts support plain paragraphs only")
                paragraph = document.add_paragraph()
                inline(paragraph, quote["tokens"])
                shade = OxmlElement("w:shd"); shade.set(qn("w:fill"), theme["callout"])
                props = paragraph._p.get_or_add_pPr()
                props.append(shade)
                borders = OxmlElement("w:pBdr")
                left_only = theme["callout_left_only"]
                for edge in (("left",) if left_only else ("top", "left", "bottom", "right")):
                    border = OxmlElement("w:" + edge)
                    for key, value in (("val", "single"), ("sz", "24" if left_only else "8"),
                                       ("color", theme["callout_border"]), ("space", "8" if left_only else "10")):
                        border.set(qn("w:" + key), value)
                    borders.append(border)
                props.append(borders)
                paragraph.paragraph_format.left_indent = Pt(6 if left_only else 8)
                paragraph.paragraph_format.right_indent = Pt(6 if left_only else 8)
                paragraph.paragraph_format.line_spacing = 1.15 if left_only else 1.12
                paragraph.paragraph_format.space_before = paragraph.paragraph_format.space_after = Pt(12)
        elif kind == "list":
            start = int(item["start"]) if item["ordered"] else 0
            for index, entry in enumerate(item["items"]):
                if entry["task"] or len(entry["tokens"]) != 1 or entry["tokens"][0]["type"] not in {"text", "paragraph"}:
                    raise ValueError("Only flat single-paragraph lists are supported")
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.left_indent = Cm(0.5)
                paragraph.paragraph_format.first_line_indent = Cm(-0.5)
                paragraph.add_run(f"{start + index}. " if item["ordered"] else "• ")
                inline(paragraph, entry["tokens"][0]["tokens"])
        else:
            raise ValueError(f"Unsupported block Markdown: {kind}; use headings, paragraphs, flat lists, tables or local images")
    data = BytesIO()
    document.save(data)
    with output_path.open("xb") as output:
        try:
            output.write(data.getvalue())
        except OSError:
            output.close()
            output_path.unlink()
            raise
    return output_path


def build_fixed(input_path: Path, output_path: Path, *, template_path: Path | None = None,
                visual_dir: Path | None = None) -> Path:
    """Render validated JSON directly; no Markdown conversion can drop evidence."""
    if output_path.exists() or output_path.is_symlink():
        raise ValueError(f"Output must be a new file: {output_path}")
    if input_path.suffix.lower() != ".json" or output_path.suffix.lower() != ".docx":
        raise ValueError("Fixed-template input must be .json and output must be .docx")
    plugin_scripts = Path(__file__).resolve().parents[3] / "scripts"
    sys.path.insert(0, str(plugin_scripts))
    from proposal_contract import image_info, load_template
    from proposal_docs import specification, table_column_proportions

    data = json.loads(input_path.read_text(encoding="utf-8-sig"))
    template = load_template(template_path)
    plan = specification(data, template)
    config = template["style"]
    document = Document()
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    page = document.sections[0]
    page.page_width, page.page_height = Pt(config["page_width_pt"]), Pt(config["page_height_pt"])
    page.top_margin = page.bottom_margin = page.left_margin = page.right_margin = Pt(config["margin_pt"])
    page.header_distance = page.footer_distance = Pt(24)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3"):
        style = document.styles[name]
        style.font.name = config["font_family"]
        style.font.color.rgb = RGBColor.from_string("17212B")
        style.font.size = Pt(config["body_pt"])
        style.paragraph_format.line_spacing = 1.1
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_together = True
        if name != "Normal":
            style.font.bold = True
            style.paragraph_format.keep_with_next = True
    footer = page.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)

    def text_run(paragraph, text, *, size=None, bold=False, link=None, muted=False):
        holder = Paragraph(OxmlElement("w:p"), paragraph._parent) if link else paragraph
        run = holder.add_run(text)
        run.font.name = config["font_family"]
        run.font.size = Pt(config["body_pt"] if size is None else size)
        run.font.bold, run.font.italic = bold, False
        run.font.color.rgb = RGBColor.from_string("236B9C" if link else "52606D" if muted else "17212B")
        if link:
            run.font.underline = True
            hyperlink = OxmlElement("w:hyperlink")
            hyperlink.set(qn("r:id"), paragraph.part.relate_to(link, RT.HYPERLINK, is_external=True))
            hyperlink.append(run._r)
            paragraph._p.append(hyperlink)
        return run

    def paragraph_block(block):
        paragraph = document.add_paragraph()
        role = block["role"]
        if role in {"table_title", "detail_title"}:
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.space_before = Pt(10)
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.paragraph_format.page_break_before = bool(block.get("page_break_before"))
        elif role in {"intro", "link", "caption"}:
            paragraph.paragraph_format.keep_with_next = True
        if role == "takeaway":
            properties = paragraph._p.get_or_add_pPr()
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "D9EAD3")
            properties.append(shade)
            borders = OxmlElement("w:pBdr")
            for edge in ("top", "left", "bottom", "right"):
                border = OxmlElement("w:" + edge)
                for key, value in (("val", "single"), ("sz", "6"), ("color", "6DA855"), ("space", "6")):
                    border.set(qn("w:" + key), value)
                borders.append(border)
            properties.append(borders)
            paragraph.paragraph_format.space_before = Pt(8)
            paragraph.paragraph_format.space_after = Pt(10)
            paragraph.paragraph_format.keep_with_next = True
        text_run(paragraph, block["text"],
                 size=config["caption_pt"] if role in {"caption", "note"} else config["body_pt"],
                 bold=role in {"takeaway", "table_title", "detail_title"},
                 link=block.get("url"), muted=role in {"caption", "note"})

    def table_block(spec):
        cols = len(spec["rows"][0])
        table = document.add_table(rows=0, cols=cols)
        table.autofit = False
        proportions = table_column_proportions(spec["key"], cols)
        for column, fraction in zip(table.columns, proportions):
            column.width = Pt(config["content_width_pt"] * fraction)
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            border = OxmlElement("w:" + edge)
            for key, value in (("val", "single"), ("sz", "4"), ("color", "D8E1E8")):
                border.set(qn("w:" + key), value)
            borders.append(border)
        table._tbl.tblPr.append(borders)
        for ri, values in enumerate(spec["rows"]):
            row = table.add_row()
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            if ri == 0 and spec["kind"] == "data":
                row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
            for ci, (cell, value) in enumerate(zip(row.cells, values)):
                cell.width = table.columns[ci].width
                properties = cell._tc.get_or_add_tcPr()
                margins = OxmlElement("w:tcMar")
                for edge in ("top", "left", "bottom", "right"):
                    margin = OxmlElement("w:" + edge)
                    margin.set(qn("w:w"), "80" if spec["section_id"] == "budget" else "100")
                    margin.set(qn("w:type"), "dxa")
                    margins.append(margin)
                properties.append(margins)
                if ri == 0:
                    shade = OxmlElement("w:shd")
                    shade.set(qn("w:fill"), "FFF2CC" if spec["kind"] == "scores" else "D9EAF7")
                    properties.append(shade)
                for pi, line in enumerate(str(value).split("\n")):
                    paragraph = cell.paragraphs[0] if pi == 0 else cell.add_paragraph()
                    paragraph.paragraph_format.space_after = Pt(1)
                    paragraph.paragraph_format.line_spacing = 1.05
                    paragraph.paragraph_format.keep_together = True
                    paragraph.paragraph_format.keep_with_next = spec["kind"] == "scores" and pi < 2
                    if spec["key"] == "roadmap:roadmap" and ci > 0:
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    text_run(paragraph, line,
                             size=20 if spec["kind"] == "scores" and pi == 1 else config["table_pt"],
                             bold=pi < 2 if spec["kind"] == "scores" else ri == 0,
                             link=spec.get("cell_links", {}).get(f"{ri}:{ci}"))

    def image_block(block):
        visual = block["visual"]
        if visual["type"] == "image":
            path = Path(visual["path"]).resolve()
        elif visual_dir is not None:
            path = (visual_dir / f"visual-{block['section_id']}.png").resolve()
        else:
            raise ValueError("JSON charts require --visual-dir with pre-rendered PNG files")
        _, width, height = image_info(path)
        available_height = config["page_height_pt"] - 2 * config["margin_pt"] - 80
        cap = available_height if visual["type"] == "image" else 180
        display_width = min(config["content_width_pt"], cap * width / height)
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.keep_with_next = True
        picture = paragraph.add_run().add_picture(str(path), width=Pt(display_width))
        picture._inline.docPr.set("descr", visual.get("alt") or
                                  " · ".join(visual.get("labels", [])))
        picture._inline.docPr.set("name", block["key"])

    for section in plan["sections"]:
        level = min(max(section["level"], 1), 3)
        heading = document.add_paragraph(style="Title" if section["id"] == "cover" else f"Heading {level}")
        heading.paragraph_format.page_break_before = section["id"] != "cover"
        heading.paragraph_format.space_after = Pt(12)
        text_run(heading, section["heading"], size=config["heading_sizes_pt"][level - 1], bold=True)
        if data["mode"] == "demo":
            badge = document.add_paragraph()
            badge.paragraph_format.keep_with_next = True
            text_run(badge, "ДЕМОНСТРАЦИОННЫЙ ПРИМЕР · СИНТЕТИЧЕСКИЕ ДАННЫЕ", size=8.5, muted=True)
        if section["id"] == "cover":
            client = data["client"]
            text_run(document.add_paragraph(), f"{client['name']} · {client['domain']}")
            text_run(document.add_paragraph(), client["date"])
        for block in section["blocks"]:
            if block["type"] == "paragraph":
                paragraph_block(block)
            elif block["type"] == "table":
                table_block(block["table"])
            elif block["type"] == "image":
                image_block(block)
    buffer = BytesIO()
    document.save(buffer)
    with output_path.open("xb") as output:
        try:
            output.write(buffer.getvalue())
        except OSError:
            output.close()
            output_path.unlink()
            raise
    return output_path


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--style", choices=STYLES, default="woodcort")
    parser.add_argument("--template", type=Path, help="Fixed-template JSON route only")
    parser.add_argument("--visual-dir", type=Path, help="Existing rendered chart PNGs for the JSON route")
    args = parser.parse_args()
    try:
        if args.input.suffix.lower() == ".json":
            output = build_fixed(args.input, args.output, template_path=args.template, visual_dir=args.visual_dir)
        elif args.template or args.visual_dir:
            raise ValueError("--template and --visual-dir apply only to fixed-template JSON")
        else:
            output = build(args.input, args.output, args.style)
    except (OSError, ValueError, subprocess.SubprocessError) as error:
        parser.exit(1, f"DOCX build failed: {error}\n")
    print(output.resolve())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
