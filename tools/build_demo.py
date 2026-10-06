"""Rebuild the synthetic reference from Markdown and CSV; never calls APIs."""
from pathlib import Path
import csv
import io
import re
import shutil
import argparse
import os

ROOT = Path(__file__).resolve().parents[1]
DEMO = ROOT / "examples/demo"


def forecast_rows():
    rows = []
    for scenario, uplift, new_visits in (("осторожный", .2, 30), ("рабочий", .3, 50), ("ускоренный", .4, 70)):
        mature = 160 * uplift + new_visits
        cumulative = 0
        for month, maturity in enumerate((0, .25, .5, 1, 1, 1), 1):
            delta = mature * maturity
            cumulative += delta
            rows.append({"scenario": scenario, "period": month, "baseline_visits": 575,
                         "additional_visits": round(delta, 2), "total_visits": round(575 + delta, 2),
                         "cumulative_additional_visits": round(cumulative, 2),
                         "assumption": "one cohort; constant baseline; synthetic inputs"})
    return rows


def write_forecast():
    rows = forecast_rows()
    with (DEMO / "forecast.csv").open("w", encoding="utf-8", newline="") as out:
        writer = csv.DictWriter(out, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def create_document():
    from PIL import Image, ImageDraw, ImageFont
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

    def font(size):
        for candidate in (os.environ.get("OFFERIST_EXAMPLE_FONT"), "arial.ttf", "DejaVuSans.ttf", "LiberationSans-Regular.ttf"):
            if candidate:
                try:
                    return ImageFont.truetype(candidate, size)
                except OSError:
                    pass
        raise ValueError("Set OFFERIST_EXAMPLE_FONT to a TrueType font with Cyrillic support")
    with (DEMO / "pages.csv").open(encoding="utf-8") as source:
        pages = list(csv.DictReader(source))
    traffic = [sum(int(p[engine + "_visits"]) for p in pages) for engine in ("yandex", "google")]
    groups = [sum(int(p["yandex_visits"]) + int(p["google_visits"]) for p in pages if p["group"] == group)
              for group in ("strong", "quick", "weak", "not_applicable")]
    for name, labels, values, colors in (
        ("traffic", ["Яндекс", "Google"], traffic, ["#3D89DB", "#58AD78"]),
        ("groups", ["Сильные", "Быстрый рост", "Слабые", "Сервисная"], groups, ["#3D89DB", "#58AD78", "#A7B2BC", "#C4CED7"])):
        chart = Image.new("RGB", (1400, 520), "white")
        draw = ImageDraw.Draw(chart)
        draw.text((85, 22), "Учебный сентябрь 2026 · посещения", font=font(28), fill="#333333")
        ceiling = max(values) * 1.2
        for step in range(5):
            y = 430 - step * 80
            draw.line((85, y, 1350, y), fill="#E5E9EC", width=2)
            draw.text((10, y - 12), str(round(ceiling * step / 4)), font=font(22), fill="#666666")
        column = 1200 / len(values)
        for index, (label, value, color) in enumerate(zip(labels, values, colors)):
            x = 115 + column * index
            height = value / ceiling * 320
            if value:
                draw.rectangle((x + column * .22, 430 - height, x + column * .78, 430), fill=color)
            draw.text((x + column / 2, 410 - height), str(value), font=font(26), fill="#333333", anchor="mm")
            draw.text((x + column / 2, 464), label, font=font(24), fill="#333333", anchor="mm")
        chart.save(DEMO / (name + ".png"), dpi=(180, 180))

    document = Document()
    document.core_properties.title = "Учебное коммерческое предложение для мебельного проекта"
    document.core_properties.author = "SEOMEAT"
    document.core_properties.subject = "Synthetic example; no client data"
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Cm(2.29)
    for name, size in (("Normal", 11.5), ("Title", 30), ("Heading 1", 24), ("Heading 2", 16), ("Heading 3", 13)):
        style = document.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(10)
        for border in style.element.xpath(".//w:pBdr"):
            border.getparent().remove(border)
        style.font.underline = False
    document.styles["Normal"].paragraph_format.line_spacing = 1.08
    footer = section.footer.paragraphs[0]
    footer.add_run("Учебный пример · example.org · ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    for run in footer.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string("777777")

    lines = (DEMO / "proposal.md").read_text(encoding="utf-8").splitlines()
    title_seen = False
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("# "):
            if title_seen:
                document.add_page_break()
                document.add_heading(line[2:], level=1)
            else:
                document.add_paragraph(line[2:], "Title")
                title_seen = True
        elif line.startswith("## "):
            document.add_heading(line[3:], level=2)
        elif line.startswith("!["):
            image = re.fullmatch(r"!\[(.*?)\]\((.*?)\)", line)
            document.add_picture(str(DEMO / image.group(2)), width=Cm(16.2))
            caption = document.add_paragraph(image.group(1))
            caption.paragraph_format.space_after = Pt(8)
            for run in caption.runs:
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor.from_string("666666")
        elif line.startswith("|"):
            raw = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                row = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r"[-: ]+", c) for c in row):
                    raw.append(row)
                i += 1
            table = document.add_table(rows=0, cols=len(raw[0]))
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False
            widths = {3: [6, 5.2, 5.2], 4: [4.2, 4.1, 4.1, 4.0]}[len(raw[0])]
            for column, width in zip(table.columns, widths):
                column.width = Cm(width)
            for row_number, cells in enumerate(raw):
                row = table.add_row()
                if row_number == 0:
                    repeat = OxmlElement("w:tblHeader")
                    row._tr.get_or_add_trPr().append(repeat)
                for col, text in enumerate(cells):
                    cell = row.cells[col]
                    cell.width = Cm(widths[col])
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    cell.text = text
                    props = cell._tc.get_or_add_tcPr()
                    shade = OxmlElement("w:shd")
                    shade.set(qn("w:fill"), "DCECF7" if row_number == 0 else ("F5F7F8" if row_number % 2 == 0 else "FFFFFF"))
                    props.append(shade)
                    borders = OxmlElement("w:tcBorders")
                    for edge in ("top", "left", "bottom", "right"):
                        border = OxmlElement("w:" + edge)
                        for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                            border.set(qn("w:" + key), value)
                        borders.append(border)
                    props.append(borders)
                    margins = OxmlElement("w:tcMar")
                    for edge, value in (("top", "90"), ("bottom", "90"), ("left", "100"), ("right", "100")):
                        margin = OxmlElement("w:" + edge)
                        margin.set(qn("w:w"), value)
                        margin.set(qn("w:type"), "dxa")
                        margins.append(margin)
                    props.append(margins)
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.space_after = Pt(2)
                        paragraph.paragraph_format.line_spacing = 1.05
                        for run in paragraph.runs:
                            run.font.size = Pt(10)
                            run.bold = row_number == 0
            document.add_paragraph()
            continue
        else:
            document.add_paragraph(line)
        i += 1
    output = DEMO / "proposal.docx"
    for paragraph in document.paragraphs:
        for border in paragraph._p.xpath(".//w:pBdr"):
            border.getparent().remove(border)
    document.save(output)
    shutil.copyfile(output, ROOT / "plugins/offerist/skills/offerist-quick/assets/reference-proposal.docx")
    print(str(output))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--forecast-only", action="store_true")
    args = parser.parse_args()
    write_forecast()
    if not args.forecast_only:
        create_document()


if __name__ == "__main__":
    main()
